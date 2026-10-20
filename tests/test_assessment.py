"""Тесты формирования перечня мер, отчёта о несоответствиях и метрик."""

import pytest

from engine.assessment import assess, to_markdown
from engine.dataset import implemented_codes, load_objects
from engine.measures import extract_codes, required_measures
from engine.report_docx import to_docx
from engine.metrics import Counts, category_accuracy
from engine.rules_loader import load_rules


@pytest.fixture(scope="module")
def rules():
    return load_rules()


def codes(measures):
    return {m["code"] for m in measures}


def test_required_by_category(rules):
    assert len(required_measures(rules, 3)) == 87
    assert len(required_measures(rules, 2)) == 110
    assert len(required_measures(rules, 1)) == 124
    assert required_measures(rules, None) == []


def test_specific_measures(rules):
    # Примеры из приложения к приказу № 239
    assert "УПД.3" not in codes(required_measures(rules, 3))   # доверенная загрузка — с категории 2
    assert "УПД.3" in codes(required_measures(rules, 2))
    assert "АУД.5" in codes(required_measures(rules, 1))       # только для категории 1
    assert "АУД.5" not in codes(required_measures(rules, 2))
    assert "ИАФ.6" not in codes(required_measures(rules, 1))   # нет ни в одном базовом наборе


def test_dns5_literal(rules):
    # ДНС.5 по тексту приложения: категории 3 и 1, но не 2
    assert "ДНС.5" in codes(required_measures(rules, 3))
    assert "ДНС.5" not in codes(required_measures(rules, 2))
    assert "ДНС.5" in codes(required_measures(rules, 1))


def test_missing_measures(rules):
    implemented = [c for c in codes(required_measures(rules, 3)) if c not in ("АУД.2", "ИНЦ.1")]
    rep = assess(rules, "тест", "healthcare", {"1": 30}, implemented)
    assert rep.categorization.category == 3
    assert set(rep.missing_codes) == {"АУД.2", "ИНЦ.1"}
    assert rep.compliance_percent == round(100 * 85 / 87, 1)


def test_extra_measures_do_not_matter(rules):
    implemented = list(codes(required_measures(rules, 3))) + ["ЗИС.17"]
    rep = assess(rules, "тест", "healthcare", {"1": 30}, implemented)
    assert rep.missing == []
    assert rep.compliance_percent == 100.0


def test_unknown_measure_code_rejected(rules):
    with pytest.raises(KeyError):
        assess(rules, "тест", "healthcare", {"1": 30}, ["XXX.1"])


def test_no_category_no_requirements(rules):
    rep = assess(rules, "тест", "science", {}, [])
    assert rep.required == [] and rep.missing == []
    assert rep.compliance_percent is None


def test_non_monotonic_warning_for_category_2(rules):
    rep = assess(rules, "тест", "healthcare", {"1": 100}, [])
    assert any("ДНС.5" in w for w in rep.warnings)


def test_markdown_report(rules):
    rep = assess(rules, "Объект Х", "healthcare", {"1": 30}, [])
    md = to_markdown(rules, rep)
    assert "Объект Х" in md
    assert "Категория значимости: 3" in md
    assert "Несоответствий: 87" in md


def test_metrics():
    c = Counts()
    c.add({"A", "B", "C"}, {"A", "B", "D"})
    assert (c.tp, c.fp, c.fn) == (2, 1, 1)
    assert c.recall == pytest.approx(2 / 3)
    assert c.precision == pytest.approx(2 / 3)
    assert Counts().recall == 1.0 and Counts().precision == 1.0
    assert category_accuracy([(1, 1), (2, 3), (None, None), (3, 3)]) == 0.75


def test_extract_codes_from_text(rules):
    text = "Реализованы: ИАФ.1; УПД.3, AУД.2 (латинская A), 0ДТ.2, ИАФ.1 повтор, XYZ.1, ИАФ.99"
    found, unknown = extract_codes(rules, text)
    assert found == ["ИАФ.1", "УПД.3", "АУД.2", "ОДТ.2"]
    assert "ИАФ.99" in unknown


def test_docx_report(rules):
    import io

    import docx
    rep = assess(rules, "Объект Х", "healthcare", {"1": 30}, ["ИАФ.0"])
    data = to_docx(rules, rep, "healthcare", {"1": 30})
    d = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    cells = "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert "Присвоена III категория значимости" in text
    assert "Объект Х" in cells and "Здравоохранение" in cells
    assert "ИАФ.1" in cells and "ИАФ.0" not in cells.split("Перечень")[-1].split("\n")[0]
    assert "30 человек" in cells


@pytest.mark.parametrize("obj", load_objects(), ids=lambda o: o["id"])
def test_reference_objects(rules, obj):
    """Каждый объект тестового набора совпадает с эталоном (ручной расчёт)."""
    rep = assess(rules, obj["name"], obj["sphere"], obj.get("indicators") or {},
                 implemented_codes(rules, obj["implemented"]))
    assert rep.categorization.category == obj["expected_category"]
    assert set(rep.missing_codes) == set(obj["expected_missing"])
