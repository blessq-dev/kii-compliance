"""Тесты возможностей: адаптация набора мер, компенсирующие меры, ГИС/ИСПДн (п. 24),
требования приказа № 235, черновик акта, карточка эксперимента, перечень объектов, журнал эксперимента."""

import io
import re
from pathlib import Path

import docx
import pytest

from conftest import ORIGINAL_DATE
from engine import registry, trials
from engine.assessment import COMPENSATED, NOT_APPLICABLE, assess, to_markdown
from engine.measures import required_measures
from engine.report_docx import card_docx, to_act_docx, to_docx
from engine.rules_loader import load_rules
from engine.security_system import assess_security
from engine.validation import InputError

ROOT = Path(__file__).resolve().parent.parent
RULES = load_rules()
BASE3 = [m["code"] for m in required_measures(RULES, 3)]


def docx_text(data: bytes) -> str:
    d = docx.Document(io.BytesIO(data))
    return "\n".join([p.text for p in d.paragraphs] +
                     [c.text for t in d.tables for row in t.rows for c in row.cells])


def assess3(implemented, **kw):
    return assess(RULES, "Объект", "healthcare", {"1": 30}, implemented, **kw)


# ---------------- адаптация и компенсирующие меры (приказ № 239, пп. 23, 26) ----------------

def test_not_applicable_excluded_from_base():
    impl = [c for c in BASE3 if c not in ("ЗИС.32", "АУД.2")]
    rep = assess3(impl, statuses={"ЗИС.32": {"status": NOT_APPLICABLE, "text": "беспроводных сетей нет"}})
    assert rep.missing_codes == ["АУД.2"]
    assert [i["measure"]["code"] for i in rep.not_applicable] == ["ЗИС.32"]
    assert rep.applicable_count == 89
    assert rep.compliance_percent == round(100 * 88 / 89, 1)


def test_compensated_counts_as_covered():
    impl = [c for c in BASE3 if c != "АВЗ.2"]
    rep = assess3(impl, statuses={"АВЗ.2": {"status": COMPENSATED, "text": "фильтрация почты на шлюзе"}})
    assert rep.missing == [] and rep.compliance_percent == 100.0
    assert rep.compensated[0]["text"] == "фильтрация почты на шлюзе"


@pytest.mark.parametrize("statuses, fragment", [
    ({"ЗИС.32": {"status": NOT_APPLICABLE, "text": "  "}}, "обоснование"),
    ({"АВЗ.2": {"status": COMPENSATED, "text": ""}}, "описание компенсирующей меры"),
    ({"ИАФ.0": {"status": "забыли", "text": "x"}}, "недопустимый статус"),
    ({"XXX.1": {"status": NOT_APPLICABLE, "text": "x"}}, "неизвестная мера"),
])
def test_invalid_statuses_rejected(statuses, fragment):
    with pytest.raises(InputError, match=fragment):
        assess3([], statuses=statuses)


def test_status_conflicts_with_implemented():
    with pytest.raises(InputError, match="одновременно"):
        assess3(["ИАФ.0"], statuses={"ИАФ.0": {"status": NOT_APPLICABLE, "text": "x"}})


def test_details_due_date_format():
    with pytest.raises(InputError, match="ГГГГ-ММ-ДД"):
        assess3([], details={"ИАФ.0": {"due": "01.12.2027"}})
    rep = assess3([], details={"ИАФ.0": {"due": "2027-12-01"}, "ЗИС.17": {"means": "x"}})
    assert "ИАФ.0" in rep.details and "ЗИС.17" not in rep.details  # ЗИС.17 вне базового набора


# ---------------- ГИС / ИСПДн (п. 24) ----------------

def test_gis_ispdn_notice():
    rep = assess3([], attributes={"is_gis": True, "is_ispdn": True})
    assert len(rep.notices) == 1 and "п. 24" in rep.notices[0]
    assert "государственной информационной системой и информационной системой персональных данных" in rep.notices[0]
    assert assess3([]).notices == []
    notices = assess(RULES, "x", "science", {}, [], attributes={"is_gis": True}).notices  # не значимый
    assert not any("п. 24" in n for n in notices)


# ---------------- приказ № 235 ----------------

def test_security_rules_structure():
    reqs = RULES.security["requirements"]
    assert len(reqs) == 48
    assert all("№ 235, п." in r["source"] for r in reqs)
    assert {r["id"] for r in reqs if r.get("condition")} == {"С.3", "С.11", "С.14", "С.15", "С.17", "С.18",
                                                             "Т.6", "Т.8", "Д.7", "Ф.12"}
    assert len({r["id"] for r in reqs}) == 48


def test_security_rules_match_notes():
    """Чек-лист выжимки = пакет действующей редакции; таблица отличий + чек-лист = пакет редакции 2019 г."""
    text = (ROOT / "notes" / "fstek235_requirements.md").read_text(encoding="utf-8")
    checklist, diff = text.split("## Редакция 2019 г.")
    rows = re.findall(r"^\| ([СТДФ]\.\d+) \| (.+?) \| (.+?) \|", checklist, re.M)
    assert [(i, t) for i, t, _ in rows] == [(r["id"], r["text"]) for r in RULES.security["requirements"]]
    old_texts = dict((i, t) for i, t, _ in re.findall(r"^\| ([СТДФ]\.\d+) \| (.+?) \| (.+?) \|$", diff, re.M))
    added = {"С.18", "Т.8"}
    expected = [(i, old_texts.get(i, t)) for i, t, _ in rows if i not in added]
    original = load_rules(on_date=ORIGINAL_DATE).security["requirements"]
    assert expected == [(r["id"], r["text"]) for r in original]


def test_assess_security():
    s = assess_security(RULES, ["С.1", "Д.1"], ["С.3"])
    assert [r["id"] for r in s.done] == ["С.1", "Д.1"]
    assert s.applicable_count == 47 and len(s.missing) == 45
    assert s.percent == round(100 * 2 / 47, 1)


@pytest.mark.parametrize("done, na, fragment", [
    (["Ж.1"], [], "неизвестное требование"),
    ([], ["С.1"], "не является условным"),
    (["С.3"], ["С.3"], "и как выполненное"),
])
def test_assess_security_errors(done, na, fragment):
    with pytest.raises(InputError, match=fragment):
        assess_security(RULES, done, na)


def test_security_only_for_significant_objects():
    assert assess3([], security={"done": ["С.1"]}).security.done[0]["id"] == "С.1"
    assert assess(RULES, "x", "science", {}, [], security={"done": ["С.1"]}).security is None


# ---------------- отчёты ----------------

def full_report():
    impl = [c for c in BASE3 if c not in ("ЗИС.32", "АВЗ.2", "АУД.2")]
    return assess3(impl,
                   statuses={"ЗИС.32": {"status": NOT_APPLICABLE, "text": "беспроводных сетей нет"},
                             "АВЗ.2": {"status": COMPENSATED, "text": "фильтрация почты на шлюзе"}},
                   details={"АУД.2": {"due": "2027-06-30"}, "ИАФ.0": {"means": "Политика ИБ, приказ № 15"}},
                   attributes={"is_ispdn": True}, security={"done": ["С.1"], "not_applicable": ["С.3"]})


def test_markdown_contains_new_sections():
    md = to_markdown(RULES, full_report())
    for fragment in ("Меры, исключённые при адаптации", "беспроводных сетей нет", "Компенсирующие меры",
                     "2027-06-30", "приказ ФСТЭК № 235", "п. 24", "Документы базы правил",
                     "в ред. от 28.08.2024"):
        assert fragment in md, fragment


def test_docx_report_contains_new_sections():
    text = docx_text(to_docx(RULES, full_report(), "healthcare", {"1": 30}, {"is_ispdn": True}))
    for fragment in ("исключённые при адаптации", "фильтрация почты на шлюзе", "30.06.2027",
                     "Политика ИБ, приказ № 15", "приказ ФСТЭК России № 235", "информационная система персональных данных"):
        assert fragment in text, fragment


def test_act_docx_follows_p16():
    text = docx_text(to_act_docx(RULES, full_report(), "healthcare", {"1": 30}, {"is_ispdn": True}))
    for fragment in ("АКТ", "п. 16 Правил категорирования", "1. Сведения об объекте",
                     "2. Результаты анализа угроз", "3. Сведения о присвоенной категории",
                     "4. Реализованные меры", "5. Сведения о необходимых мерах",
                     "Присвоена III категория значимости", "(п. 11 «в»)", "ЧЕРНОВИК"):
        assert fragment in text, fragment


def test_card_has_no_answers():
    obj = {"id": "obj99", "name": "Тест", "sphere": "healthcare", "indicators": {"1": 30, "2а": "mo"}}
    text = docx_text(card_docx(RULES, obj, ["ИАФ.0", "ИАФ.1"]))
    assert "30 человек" in text and "В пределах территории одного муниципального" in text and "ИАФ.0, ИАФ.1" in text
    assert "категория" not in text.lower().replace("категорию значимости", "")  # без ответа


# ---------------- перечень объектов ----------------

def test_registry_roundtrip(tmp_path):
    state = {"object_name": "МИС больницы № 1", "sphere": "healthcare", "indicators": {"1": 30}}
    p = registry.save(state, {"category": 3, "compliance": 90.0, "missing": 5}, tmp_path)
    registry.save({**state, "object_name": "АСУ ТП котельной"}, {"category": 2}, tmp_path)
    (tmp_path / "битый.json").write_text("{не json", encoding="utf-8")
    items = registry.list_all(tmp_path)
    assert {i["state"]["object_name"] for i in items} == {"МИС больницы № 1", "АСУ ТП котельной"}
    assert registry.load(p.name, tmp_path) == state
    registry.save({**state, "indicators": {"1": 60}}, {"category": 2}, tmp_path)  # перезапись того же объекта
    assert registry.load(p.name, tmp_path)["indicators"] == {"1": 60}
    registry.delete(p.name, tmp_path)
    assert len(registry.list_all(tmp_path)) == 1


# ---------------- журнал эксперимента ----------------

def test_trials_compare_and_log(tmp_path):
    log = tmp_path / "log.csv"
    ok = trials.append(log, "Иванов", "obj02", 4.5, 2, ["СОВ.1", "ЗИС.4"])
    bad = trials.append(log, "Петров", "obj02", 12, 3, ["СОВ.1", "ИАФ.0"])
    assert ok["category_ok"] == 1 and (ok["tp"], ok["fp"], ok["fn"]) == (2, 0, 0)
    assert bad["category_ok"] == 0 and (bad["tp"], bad["fp"], bad["fn"]) == (1, 1, 1)
    rows = trials.read(log)
    assert [r["participant"] for r in rows] == ["Иванов", "Петров"] and rows[1]["minutes"] == "12.0"
    assert trials.compare("obj05", None, [])["category_ok"]  # «без категории» тоже сравнивается



# ---------------- проверки адаптации и пустые показатели ----------------

def test_all_measures_excluded_compliance_not_defined():
    statuses = {c: {"status": NOT_APPLICABLE, "text": "технология не используется"} for c in BASE3}
    rep = assess3([], statuses=statuses)
    assert rep.compliance_percent is None
    assert rep.compliance_text.startswith("не определён")
    assert "None" not in to_markdown(RULES, rep)
    assert any("Все меры базового набора исключены" in n for n in rep.notices)


@pytest.mark.parametrize("text, fragment", [
    ("ИАФ.1", "словами"),
    (". . .", "словами"),
    ("ИАФ.1 компенсирована ИАФ.1", "не может совпадать"),
    ("компенсирована мерой ЖЖЖ.7", "неизвестный код"),
])
def test_compensation_text_checked(text, fragment):
    with pytest.raises(InputError, match=fragment):
        assess3([], statuses={"ИАФ.1": {"status": COMPENSATED, "text": text}})


def test_compensation_by_unimplemented_measure_noted():
    rep = assess3([], statuses={"ИАФ.1": {"status": COMPENSATED, "text": "заменена мерой ИАФ.2 и пропусками"}})
    assert any("компенсирована мерой ИАФ.2" in n for n in rep.notices)
    assert any("п. 26" in n for n in rep.notices)
    rep = assess3(["ИАФ.2"], statuses={"ИАФ.1": {"status": COMPENSATED, "text": "заменена мерой ИАФ.2"}})
    assert not any("компенсирована мерой" in n for n in rep.notices)


def test_policy_excluded_while_group_remains_noted():
    rep = assess3([], statuses={"ИАФ.0": {"status": NOT_APPLICABLE, "text": "политика не нужна"}})
    assert any("ИАФ.0 (правила и процедуры группы)" in n for n in rep.notices)


def test_unfilled_indicators_listed_and_noted():
    rep = assess(RULES, "x", "healthcare", {}, [])
    assert any("Ни один показатель" in n for n in rep.notices)
    assert any("Не заполнены и считаются неприменимыми" in e for e in rep.categorization.explanation)
    rep = assess(RULES, "x", "healthcare", {"1": 30}, [])
    assert not any("Ни один показатель" in n for n in rep.notices)
    assert "1" not in rep.categorization.not_filled and len(rep.categorization.not_filled) == 29
