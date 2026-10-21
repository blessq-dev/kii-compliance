"""Выгрузка документов в формате DOCX.

  to_docx()      — отчёт об оценке соответствия (для приложения к акту);
  to_act_docx()  — черновик акта категорирования по составу сведений п. 16 Правил (ПП РФ № 127);
  card_docx()    — карточка тестового объекта для участника эксперимента.

Модуль только оформляет готовые результаты: правил и вычислений категории здесь нет.
"""

import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from engine.assessment import AssessmentReport
from engine.measures import group_title
from engine.rules_loader import RuleBase

ROMAN = {1: "I", 2: "II", 3: "III"}
GREY = "555555"


# ---------------- оформление ----------------

def _new_doc() -> Document:
    doc = Document()
    for s in doc.sections:
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(2.5), Cm(1.5), Cm(2), Cm(2)
    return doc


def _font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _p(doc, text="", size=11, bold=False, align=None, color=None, space_after=4, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if text:
        _font(p.add_run(text), size, bold, color, italic)
    return p


def _shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _table(doc, header, rows, widths):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        _font(c.paragraphs[0].add_run(h), 10, bold=True)
        _shade(c, "E8EEF8")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            _font(cells[i].paragraphs[0].add_run(str(v)), 10)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    return t


def _bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _fmt_value(rule: dict, value) -> str:
    if rule["type"] == "number":
        return f"{value:g}".replace(".", ",") + f" {rule['unit']}"
    return next(o["text"] for o in rule["options"] if o["id"] == value)


def _sphere_title(rules: RuleBase, sphere: str) -> str:
    return next((s["title"] for s in rules.spheres["spheres"] if s["id"] == sphere),
                rules.spheres["other"]["title"])


def _verdict(report: AssessmentReport) -> tuple[str, str]:
    c = report.categorization
    if not c.is_kii_subject:
        return "Объект не подлежит категорированию", GREY
    if c.category is None:
        return "Категория значимости не присваивается", GREY
    return f"Присвоена {ROMAN[c.category]} категория значимости", "B42318"


def _indicator_rows(rules, report, indicator_values):
    by_id = {r.indicator_id: r for r in report.categorization.details}
    rows = []
    for rule in rules.categorization["indicators"]:
        v = indicator_values.get(rule["id"])
        if v not in (None, ""):
            res = by_id.get(rule["id"])
            cat = ROMAN[res.category] + " категория" if res and res.category else "не соответствует"
            rows.append([rule["id"], rule["short"], _fmt_value(rule, v), cat])
    return rows


def _attributes_text(attributes: dict | None) -> str:
    a = attributes or {}
    parts = [t for t, k in (("государственная информационная система", "is_gis"),
                            ("информационная система персональных данных", "is_ispdn")) if a.get(k)]
    return ", ".join(parts) if parts else "нет"


def _conditions_text(rules: RuleBase, attributes: dict | None) -> str:
    """Подтверждённые условия применимости показателей ПП № 127."""
    a = attributes or {}
    parts = [c["short"] for c in rules.categorization["conditions"] if a.get(c["id"])]
    return "; ".join(parts) if parts else "не подтверждены"


def _measures_block(doc, rules, report, number):
    r = report
    _p(doc, f"{number}. Соответствие требованиям приказа ФСТЭК России № 239", 12, True)
    if not r.required:
        _p(doc, "Требования к значимым объектам не применяются: категория не присвоена.")
        return
    _table(doc, ["Обязательных мер", "Не применяются", "Реализовано / компенсировано", "Несоответствий",
                 "Уровень соответствия"],
           [[len(r.required), len(r.not_applicable), f"{r.covered_count - len(r.compensated)} / {len(r.compensated)}",
             len(r.missing), f"{r.compliance_percent}%".replace(".", ",")]], [3, 3, 4, 3, 4])
    _p(doc)
    if r.missing:
        _p(doc, "Перечень несоответствий (обязательные меры не реализованы):", 11, True)
        _table(doc, ["Мера", "Содержание меры", "Группа мер", "Плановый срок"],
               [[m["code"], m["title"], group_title(rules, m["group"]),
                 _date(r.details.get(m["code"], {}).get("due"))] for m in r.missing], [1.8, 7.5, 5, 2.7])
    else:
        _p(doc, "Несоответствий не выявлено.", 11, True, color="1E7B34")
    for title, items in (("Меры, исключённые при адаптации базового набора (п. 23, абз. 3):", r.not_applicable),
                         ("Компенсирующие меры (п. 26):", r.compensated)):
        if items:
            _p(doc)
            _p(doc, title, 11, True)
            _table(doc, ["Мера", "Содержание меры", "Обоснование"],
                   [[i["measure"]["code"], i["measure"]["title"], i["text"]] for i in items], [1.8, 7, 8.2])
    means = [(c, d["means"]) for c, d in r.details.items() if d.get("means")]
    if means:
        _p(doc)
        _p(doc, "Сведения о реализации мер:", 11, True)
        _table(doc, ["Мера", "Чем реализована"], means, [2, 15])


def _security_block(doc, report, number):
    s = report.security
    if not s:
        return
    _p(doc)
    _p(doc, f"{number}. Требования к системе безопасности (приказ ФСТЭК России № 235)", 12, True)
    _p(doc, f"Выполнено {len(s.done)} из {s.applicable_count} применимых требований "
            f"({str(s.percent).replace('.', ',')}%).")
    if s.missing:
        _table(doc, ["Код", "Невыполненное требование", "Источник"],
               [[q["id"], q["text"], q["source"].replace("Приказ ФСТЭК России № 235, ", "")] for q in s.missing],
               [1.4, 11.6, 4])


def _footer(doc, rules, report):
    _p(doc)
    if report.notices:
        for n in report.notices:
            _p(doc, "• " + n, 10)
    if report.warnings:
        _p(doc, "Замечания к тексту нормативных документов:", 10, True)
        for w in report.warnings:
            _p(doc, "• " + w, 9)
    _p(doc, "Документы базы правил: " + "; ".join(report.editions) + ".", 9, color=GREY)


def _date(iso: str | None) -> str:
    if not iso:
        return "—"
    y, m, d = str(iso).split("-")
    return f"{d}.{m}.{y}"


# ---------------- отчёт об оценке соответствия ----------------

def to_docx(rules: RuleBase, report: AssessmentReport, sphere: str, indicator_values: dict,
            attributes: dict | None = None) -> bytes:
    """Отчёт об оценке соответствия (байты DOCX)."""
    doc = _new_doc()
    c = report.categorization
    _p(doc, "ОТЧЁТ", 14, True, WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _p(doc, "об оценке соответствия объекта КИИ требованиям безопасности", 12, True, WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, f"Дата формирования: {datetime.now():%d.%m.%Y %H:%M}", 10, align=WD_ALIGN_PARAGRAPH.RIGHT)

    _p(doc, "1. Сведения об объекте", 12, True)
    _table(doc, ["Параметр", "Значение"],
           [["Наименование объекта", report.object_name],
            ["Сфера деятельности (187-ФЗ, ст. 2, п. 8)", _sphere_title(rules, sphere)],
            ["Является ГИС / ИСПДн", _attributes_text(attributes)],
            ["Условия применимости показателей (ПП РФ № 127)", _conditions_text(rules, attributes)]], [6, 11])

    _p(doc)
    _p(doc, "2. Значения показателей критериев значимости (ПП РФ № 127)", 12, True)
    rows = _indicator_rows(rules, report, indicator_values)
    if rows:
        _table(doc, ["№", "Показатель", "Значение", "Результат"], rows, [1.2, 6.5, 6, 3.3])
    else:
        _p(doc, "Показатели не заполнены (неприменимы к объекту).")

    _p(doc)
    _p(doc, "3. Категория значимости", 12, True)
    verdict, color = _verdict(report)
    _p(doc, verdict, 12, True, color=color)
    for line in c.explanation:
        _p(doc, "• " + line, 10)

    _p(doc)
    _measures_block(doc, rules, report, 4)
    _security_block(doc, report, 5)
    _p(doc)
    _p(doc, "Оценка выполнена по базовому набору мер с учётом отмеченной адаптации (п. 23 Требований, утв. "
            "приказом ФСТЭК России № 239). Дополнение набора мерами иных НПА не учитывается.", 9, color=GREY)
    _footer(doc, rules, report)
    _p(doc)
    _p(doc, "Оценку выполнил: ____________________ / ____________________ /", 11)
    _p(doc, "                         (подпись)                    (ФИО)", 9, color=GREY)
    return _bytes(doc)


# ---------------- черновик акта категорирования ----------------

# Состав комиссии по категорированию — ПП РФ № 127, Правила, п. 11 (подп. «а»–«д»)
COMMISSION = [
    "Председатель комиссии — руководитель субъекта КИИ или уполномоченное им лицо (п. 11 «а», п. 13)",
    "Работник — специалист в области выполняемых функций (видов деятельности) / ИТ и связи / эксплуатации "
    "основного технологического оборудования (п. 11 «б»)",
    "Работник, на которого возложены функции обеспечения безопасности объектов КИИ (п. 11 «в»)",
    "Работник подразделения по защите государственной тайны — при обработке сведений, составляющих "
    "государственную тайну (п. 11 «г»)",
    "Работник подразделения по гражданской обороне и защите от ЧС или уполномоченный на решение этих задач "
    "(п. 11 «д»)",
]


def to_act_docx(rules: RuleBase, report: AssessmentReport, sphere: str, indicator_values: dict,
                attributes: dict | None = None) -> bytes:
    """Черновик акта по результатам категорирования (байты DOCX).

    Разделы соответствуют составу сведений акта по п. 16 Правил категорирования (ПП РФ № 127).
    Сведения, которые система не формирует (результаты анализа угроз), оставлены для заполнения комиссией.
    """
    doc = _new_doc()
    _p(doc, "ЧЕРНОВИК — сформирован автоматически, подлежит проверке комиссией", 9, color="B42318",
       align=WD_ALIGN_PARAGRAPH.RIGHT)
    _p(doc, "УТВЕРЖДАЮ", 11, True, WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    _p(doc, "Руководитель субъекта КИИ (уполномоченное лицо)", 10, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    _p(doc, "____________________ / ____________________ /", 10, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    _p(doc, "«___» ____________ 20___ г.", 10, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _p(doc)
    _p(doc, "АКТ", 14, True, WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _p(doc, "по результатам категорирования объекта критической информационной инфраструктуры", 12, True,
       WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, "Оформлен в соответствии с п. 16 Правил категорирования объектов КИИ РФ, утв. постановлением "
            "Правительства РФ от 08.02.2018 № 127.", 9, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

    _p(doc, "1. Сведения об объекте критической информационной инфраструктуры", 12, True)
    _table(doc, ["Параметр", "Значение"],
           [["Наименование объекта", report.object_name],
            ["Сфера деятельности субъекта КИИ (187-ФЗ, ст. 2, п. 8)", _sphere_title(rules, sphere)],
            ["Является ГИС / ИСПДн", _attributes_text(attributes)],
            ["Условия применимости показателей (ПП РФ № 127)", _conditions_text(rules, attributes)],
            ["Назначение, архитектура, программные и программно-аппаратные средства, взаимодействие с другими "
             "объектами, доступ к сетям связи (п. 10 «а»)", "[заполняется комиссией]"]], [8, 9])

    _p(doc)
    _p(doc, "2. Результаты анализа угроз безопасности информации объекта", 12, True)
    _p(doc, "[заполняется комиссией: возможные действия нарушителей, иные источники угроз, угрозы и "
            "уязвимости, которые могут привести к компьютерным инцидентам (п. 14 «г», «д» Правил)]", 11,
       italic=True, color=GREY)

    _p(doc)
    _p(doc, "3. Сведения о присвоенной категории значимости", 12, True)
    verdict, color = _verdict(report)
    _p(doc, verdict, 12, True, color=color)
    rows = _indicator_rows(rules, report, indicator_values)
    if rows:
        _p(doc, "Результаты оценки показателей критериев значимости:", 11)
        _table(doc, ["№", "Показатель", "Значение", "Результат"], rows, [1.2, 6.5, 6, 3.3])
    for line in report.categorization.explanation:
        _p(doc, "• " + line, 10)

    _p(doc)
    _p(doc, "4. Реализованные меры по обеспечению безопасности объекта", 12, True)
    if report.required:
        implemented = [m for m in report.required if m["code"] in report.implemented]
        if implemented:
            _table(doc, ["Мера", "Содержание меры", "Чем реализована"],
                   [[m["code"], m["title"], report.details.get(m["code"], {}).get("means", "")]
                    for m in implemented], [1.8, 9, 6.2])
        else:
            _p(doc, "Реализованные меры из базового набора не отмечены.")
    else:
        _p(doc, "Категория не присвоена — меры базового набора не определяются.")

    _p(doc)
    _p(doc, "5. Сведения о необходимых мерах по обеспечению безопасности", 12, True)
    _p(doc, "В соответствии с Требованиями по обеспечению безопасности значимых объектов КИИ РФ "
            "(приказ ФСТЭК России № 239).", 10, color=GREY)
    _measures_block(doc, rules, report, "5.1")

    _p(doc)
    _p(doc, "Члены комиссии по категорированию (ПП РФ № 127, Правила, п. 11):", 11, True)
    for role in COMMISSION:
        _p(doc, role, 10, space_after=0)
        _p(doc, "____________________ / ____________________ /", 10)
    _footer(doc, rules, report)
    return _bytes(doc)


# ---------------- карточка объекта для эксперимента ----------------

def card_docx(rules: RuleBase, obj: dict, implemented: list[str]) -> bytes:
    """Карточка тестового объекта: исходные данные для участника эксперимента (без ответов)."""
    doc = _new_doc()
    _p(doc, f"Карточка объекта {obj['id']}", 14, True, WD_ALIGN_PARAGRAPH.CENTER)
    _p(doc, "Исходные данные для оценки. Определите категорию значимости (ПП РФ № 127) и перечень "
            "несоответствий — обязательных мер базового набора (приказ ФСТЭК № 239), которые не реализованы.",
       10, color=GREY)
    _table(doc, ["Параметр", "Значение"],
           [["Наименование объекта", obj["name"]],
            ["Сфера деятельности", _sphere_title(rules, obj["sphere"])],
            ["Признаки субъекта (объекта)", _conditions_text(rules, obj.get("attributes"))]], [6, 11])
    _p(doc)
    _p(doc, "Оценка последствий компьютерного инцидента (значения показателей):", 11, True)
    rows = []
    for rule in rules.categorization["indicators"]:
        v = (obj.get("indicators") or {}).get(rule["id"])
        if v not in (None, ""):
            rows.append([rule["id"], rule["short"], _fmt_value(rule, float(v) if rule["type"] == "number" else v)])
    if rows:
        _table(doc, ["№", "Показатель", "Значение"], rows, [1.2, 8, 7.8])
    else:
        _p(doc, "Ни один показатель не применим к объекту.")
    _p(doc)
    _p(doc, f"Реализованные на объекте меры защиты ({len(implemented)}):", 11, True)
    _p(doc, ", ".join(implemented) if implemented else "нет", 10)
    return _bytes(doc)
